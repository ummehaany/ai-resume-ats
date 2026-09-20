"""Shared fixtures.

IMPORTANT - what is mocked in the automated tests:
  * Groq (the LLM)        -> `mock_llm` fixture patches groq_parser._call_groq
  * Supabase (PostgREST)  -> `fake_supabase` fixture: an in-memory simulator behind httpx.MockTransport
  * spaCy / MiniLM        -> tiny fake nlp / embedder objects
Nothing here talks to a real external service. Opt-in real-service tests live in tests/integration/.
"""
import json
import os
import re
import time
import uuid

# Test configuration must be in place BEFORE any backend module is imported.
_TEST_ENV = {
    'GROQ_API_KEY': 'test-groq-key',
    'SUPABASE_URL': 'https://test-project.supabase.co',
    'SUPABASE_ANON_KEY': 'test-anon-key',
    'SUPABASE_JWT_SECRET': 'test-jwt-secret-test-jwt-secret-1234567890',
    'ANALYZE_RATE_LIMIT_PER_MINUTE': '1000',
    'ANALYZE_RATE_LIMIT_PER_DAY': '1000',
    'GLOBAL_ANALYSES_PER_DAY': '1000',
    'PDF_RATE_LIMIT_PER_MINUTE': '1000',
}
if os.getenv('RUN_INTEGRATION') == '1':
    # Real-service run (`pytest tests/integration`): keep the caller's real credentials, only relax rate limits.
    for _k, _v in _TEST_ENV.items():
        if _k.startswith(('ANALYZE_', 'GLOBAL_', 'PDF_')):
            os.environ.setdefault(_k, _v)
else:
    os.environ.update(_TEST_ENV)

import httpx
import jwt
import numpy as np
import pytest
from fastapi.testclient import TestClient

from backend.core import config
from backend.core.rate_limit import limiter

JWT_SECRET = os.environ['SUPABASE_JWT_SECRET']
USER_A = '11111111-1111-4111-8111-111111111111'
USER_B = '22222222-2222-4222-8222-222222222222'


def make_token(user_id: str = USER_A, expires_in: int = 3600, audience: str = 'authenticated', secret: str = JWT_SECRET) -> str:
    now = int(time.time())
    return jwt.encode({'sub': user_id, 'aud': audience, 'iat': now, 'exp': now + expires_in}, secret, algorithm='HS256')


def auth_header(user_id: str = USER_A) -> dict:
    return {'Authorization': f'Bearer {make_token(user_id)}'}


# ── fakes for the heavy models ───────────────────────────────────────────────
class FakeEmbedder:
    """Deterministic bag-of-words embedder: texts sharing words get a high cosine similarity."""
    dim = 128

    def encode(self, text, convert_to_tensor=False, **kwargs):
        vec = np.zeros(self.dim, dtype=float)
        for word in re.findall(r'[a-z0-9+#.]+', str(text).lower()):
            vec[hash(word) % self.dim] += 1.0
        return vec


class _FakeDoc:
    ents = []
    noun_chunks = []


class FakeNLP:
    def __call__(self, text):
        return _FakeDoc()


# ── mocked LLM ───────────────────────────────────────────────────────────────
RESUME_JSON = {
    'name': 'Jane Doe',
    'email': 'jane@example.com',
    'phone': '+1 555 123 4567',
    'linkedin': 'linkedin.com/in/janedoe',
    'github': 'github.com/janedoe',
    'professional_summary': 'Backend engineer with 3 years of experience building Python APIs and data pipelines.',
    'skills': ['Python', 'FastAPI', 'PostgreSQL', 'Docker', 'Kubernetes'],
    'experience': [{
        'job_title': 'Software Engineer', 'company': 'Acme Corp', 'start_date': 'Jan 2022', 'end_date': 'Present',
        'duration_months': 36,
        'description': 'Developed REST APIs with FastAPI serving 10K requests/day\nReduced query latency by 40% using PostgreSQL indexes\nDeployed services with Docker',
    }],
    'education': [{'degree': 'B.Sc. Computer Science', 'institution': 'State University', 'year': '2021'}],
    'certifications': [],
    'projects': [{'title': 'Resume Scorer', 'description': 'Built an ATS scoring service in Python and FastAPI', 'technologies': ['Python', 'FastAPI']}],
    'action_verbs': ['developed', 'reduced', 'deployed', 'built', 'designed', 'implemented'],
    'keywords': ['REST APIs', 'microservices', 'CI/CD', 'data pipelines', 'backend'],
}
JD_JSON = {
    'job_title': 'Senior Backend Engineer',
    'required_skills': ['Python', 'FastAPI', 'AWS', 'Terraform'],
    'preferred_skills': ['Kubernetes'],
    'experience_required': '3+ years',
    'education_required': "Bachelor's",
    'key_responsibilities': ['Build APIs'],
    'keywords': ['REST APIs', 'AWS', 'Terraform', 'microservices'],
}


@pytest.fixture
def mock_llm(monkeypatch):
    """Replace the Groq call with canned JSON. Returns the list of prompts received (one per call)."""
    from backend.services import groq_parser
    calls = []

    def fake_call(client, system_prompt, user_prompt):
        calls.append(user_prompt)
        if 'job description parser' in system_prompt:
            return json.dumps(JD_JSON)
        return json.dumps(RESUME_JSON)

    monkeypatch.setattr(groq_parser, '_get_client', lambda: object())
    monkeypatch.setattr(groq_parser, '_call_groq', fake_call)
    return calls


# ── in-memory Supabase / PostgREST simulator ────────────────────────────────
class FakeSupabase:
    """Simulates the `analyses` table incl. Row Level Security keyed on the Bearer token's `sub`.

    This proves the backend sends the *user's* token and handles responses correctly. It does NOT
    prove that the real database policies in supabase/schema.sql are correct - see tests/integration.
    """

    def __init__(self):
        self.rows = []
        self.requests = []
        self.fail_with = None   # e.g. 500 to simulate an outage

    def _uid(self, request):
        token = request.headers.get('authorization', '').removeprefix('Bearer ')
        return jwt.decode(token, options={'verify_signature': False})['sub']

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        if request.headers.get('apikey') != 'test-anon-key':
            return httpx.Response(401, json={'message': 'Invalid API key'})
        if self.fail_with:
            return httpx.Response(self.fail_with, json={'message': 'boom'})
        uid = self._uid(request)
        params = dict(request.url.params)
        visible = [r for r in self.rows if r['user_id'] == uid]          # RLS

        def eq(value):
            return value.removeprefix('eq.') if value else None

        if request.method == 'POST':
            body = json.loads(request.content)
            if body.get('user_id') != uid:                                # RLS WITH CHECK
                return httpx.Response(403, json={'code': '42501', 'message': 'new row violates row-level security policy'})
            row = {**body, 'id': str(uuid.uuid4())}
            self.rows.append(row)
            return httpx.Response(201, json=[row])

        if 'id' in params:
            visible = [r for r in visible if r['id'] == eq(params['id'])]
        if 'user_id' in params:
            visible = [r for r in visible if r['user_id'] == eq(params['user_id'])]

        if request.method == 'GET':
            visible = sorted(visible, key=lambda r: r['created_at'], reverse=True)
            offset = int(params.get('offset', 0))
            limit = int(params['limit']) if 'limit' in params else None
            visible = visible[offset:offset + limit] if limit is not None else visible[offset:]
            return httpx.Response(200, json=visible)

        if request.method == 'DELETE':
            ids = {r['id'] for r in visible}
            self.rows = [r for r in self.rows if r['id'] not in ids]
            return httpx.Response(200, json=visible)

        return httpx.Response(405)


@pytest.fixture
def fake_supabase(monkeypatch):
    from backend.database import supabase_db
    fake = FakeSupabase()
    monkeypatch.setattr(supabase_db, '_client', lambda: httpx.AsyncClient(transport=httpx.MockTransport(fake.handler)))
    return fake


# ── app / client ─────────────────────────────────────────────────────────────
@pytest.fixture(autouse=True)
def _reset_state(monkeypatch):
    from backend.api import routes
    from backend.api import auth
    from backend.services import groq_parser
    limiter.reset()
    routes._semaphores.clear()
    groq_parser._client = None
    auth._jwks_client = None
    yield


@pytest.fixture
def client(fake_supabase):
    from backend.main import app
    app.state.nlp = FakeNLP()
    app.state.embedder = FakeEmbedder()
    return TestClient(app)


# ── sample files ─────────────────────────────────────────────────────────────
RESUME_HTML = """<html><body>
<h1>Jane Doe</h1><p>jane@example.com | +1 555 123 4567 | linkedin.com/in/janedoe | github.com/janedoe</p>
<h2>Summary</h2><p>Backend engineer with 3 years of experience building Python APIs and data pipelines.</p>
<h2>Experience</h2><p>Software Engineer, Acme Corp, Jan 2022 - Present</p>
<ul><li>Developed REST APIs with FastAPI serving 10K requests/day</li>
<li>Reduced query latency by 40% using PostgreSQL indexes</li>
<li>Deployed services with Docker</li></ul>
<h2>Projects</h2><p>Resume Scorer - Built an ATS scoring service in Python and FastAPI</p>
<h2>Education</h2><p>B.Sc. Computer Science, State University, 2021</p>
<h2>Skills</h2><p>Python, FastAPI, PostgreSQL, Docker, Kubernetes</p>
</body></html>"""


@pytest.fixture(scope='session')
def resume_pdf() -> bytes:
    weasyprint = pytest.importorskip('weasyprint')
    return weasyprint.HTML(string=RESUME_HTML).write_pdf()


@pytest.fixture(scope='session')
def resume_docx() -> bytes:
    import io
    from docx import Document
    doc = Document()
    doc.add_heading('Jane Doe', 0)
    doc.add_paragraph('Backend engineer with 3 years of experience building Python APIs and data pipelines.')
    doc.add_paragraph('Developed REST APIs with FastAPI serving 10K requests/day')
    doc.add_paragraph('Skills: Python, FastAPI, PostgreSQL, Docker, Kubernetes')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
